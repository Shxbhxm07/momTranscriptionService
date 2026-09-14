"""Render the MoM response object as a Word document in the client's service-writing format.

WHY THIS EXISTS: the Kafka contract does not return minutes over HTTP. The consumer uploads a
.docx to MinIO and sends back only the bucket and object key, so a Word file — not JSON — is the
deliverable for that path. The HTTP API is unchanged and still returns JSON.

WHY THIS LAYOUT: the client writes minutes to the Joint Services Staff Duties Manual (JSSD, 2026
edition). Each rule below cites where it comes from:
  "AD"     Vol I Part 2, Appendix AD (layout of minutes of a meeting) and its explanatory notes;
  "Ch 6"   Vol I Part 2, Chapter 6 paras 9-19 (minutes of a meeting);
  "Part 1" Vol I Part 1, Chapter 2 — the "standard tenets of service writing" AD defers to for the
           page, type, spacing, numbering, headers, dates and signature block.

WHAT A RECORDING CANNOT SUPPLY. The superscription (telephone, precedence, address, file reference),
the security classification, the secretary and the distribution are never spoken in a meeting. They
come from `meta` — the Kafka job's optional `mom_meta` object — and are left out, or left for the
secretary to fill in, when absent. The classification is never inferred: Part 1 para 11.5 has an
unclassified document carry no marking at all, which is the default, and a wrong grading on a
defence document is worse than none.

ONE ITEM FOR NOW. AD records each agenda item as its own block (ITEM I, ITEM II …) ending in its
decision. llama-service does not yet say which points and decisions belong to which agenda item, so
the discussion is recorded under a single item named after the meeting. _items() is the one place
that changes when it does.
"""
import io
import logging
import math
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from typing import Any, Dict, List, Sequence, Tuple
from xml.sax.saxutils import quoteattr

import docx
import pypdfium2
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm, Pt

from config import SOFFICE_TIMEOUT

logger = logging.getLogger(__name__)
LEFT, CENTRE, JUSTIFY = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY

# ── page, type and spacing ───────────────────────────────────────────────────────────────────────
FONT, SIZE = "Arial", Pt(12)   # Part 1 para 13 names Arial but the sentence drops the size; 12 pt is
                               # what the manual, "printed to conform to the rules" (para 2), uses.
LINE_SPACING = 1.15            # Part 1 para 8.3: line spacing within a paragraph.
DRAFT_LINE_SPACING = 1.5       # Part 1 para 73: drafts are typed in one-and-a-half spacing.
EDGE = Inches(0.5)             # Part 1 paras 10.1, 10.3: top, bottom, left and right margins.
GUTTER = Inches(0.8)           # Part 1 para 10.2; AD marks the binding margin as 0.8" + 0.5".
PAGE_W, PAGE_H = Mm(210), Mm(297)              # Part 1 para 6: A4.
TEXT_W = Emu(PAGE_W - GUTTER - 2 * EDGE)       # about 6.47 in
TAB = Inches(0.5)              # Part 1 paras 33.3.6, 33.4.4: from a paragraph number to its text.
COL = Inches(1.0)              # AD: each of the Action and Info columns.
SIGNATURE_FEEDS = 9            # AD: "as required, usually nine" line feeds down to the signature.
LINES_PER_PAGE = 42            # only for the fallback page-count estimate
DOTS = "……………"                # AD marks what the secretary fills in with a row of dots.

# ── security classification ──────────────────────────────────────────────────────────────────────
# Written in full: Part 1 para 12.1 — never abbreviated in service writing.
_GRADES = {"TOP SECRET": "TOP SECRET", "TOPSEC": "TOP SECRET", "SECRET": "SECRET",
           "CONFIDENTIAL": "CONFIDENTIAL", "CONFD": "CONFIDENTIAL",
           "RESTRICTED": "RESTRICTED", "RESTD": "RESTRICTED"}
_UNGRADED = {"", "UNCLASSIFIED", "UNCLAS", "NONE", "NIL"}
_PAGES_SHOWN = {"TOP SECRET", "SECRET", "CONFIDENTIAL"}      # Part 1 App B note 2
_COPY_IN_WATERMARK = {"TOP SECRET", "SECRET"}                # Part 1 para 12.1

_EMPTY = {"", "none", "n/a", "na", "nil", "unknown", "not specified", "not stated", "not mentioned", "-"}
# "Secretary" alone is how the pipeline labels the meeting's secretary; "Secretary (Defence)" and the
# like are appointments, not the minute-taker, so only the bare word or an explicit phrase counts.
_CHAIR = re.compile(r"(the )?(chair|chairman|chairperson|chairwoman|presiding officer)", re.I)
_SECRETARY = re.compile(r"(the )?((meeting|conference) )?secretary( (of|for|to) the (meeting|conference))?|minute[s]? taker", re.I)
_WORDS = ("zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine")
_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d %B %Y", "%d %b %Y", "%d %B, %Y",
                 "%B %d, %Y", "%b %d, %Y", "%B %d %Y", "%b %d %Y", "%d %b %y", "%d-%b-%Y", "%d-%b-%y")


# ── text conventions ─────────────────────────────────────────────────────────────────────────────
def _clean(value: Any) -> str:
    s = re.sub(r"\s+", " ", str(value or "")).strip()
    return "" if s.lower().strip(".") in _EMPTY else s


def _lines(value: Any) -> List[str]:
    items = value if isinstance(value, (list, tuple)) else str(value or "").splitlines()
    return [s for s in (_clean(v) for v in items) if s]


def _grade(value: Any) -> str:
    """The document's security classification in full, or '' when it is unclassified."""
    key = re.sub(r"\s+", " ", str(value or "")).strip().upper()
    return "" if key in _UNGRADED else _GRADES.get(key, key)


def _stop(text: str) -> str:
    """Part 2 Ch 10 para 25.1: every paragraph and sub-paragraph ends with a full stop."""
    text = text.strip().rstrip(";,")
    return text if not text or text.endswith((".", "?", "!", ":-")) else text + "."


def _date(value: Any) -> str:
    """Any common way of writing a date → '11 Sep 26' (Part 1 para 40.2); '' if it can't be read."""
    s = _clean(value)
    s = re.sub(r"^(mon|tues?|wed(nes)?|thu(rs)?|fri|sat(ur)?|sun)(day)?,?\s+", "", s, flags=re.I)
    s = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", s)
    iso = re.match(r"\d{4}-\d{2}-\d{2}", s)
    for text in ([iso.group(0)] if iso else []) + [s]:
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt).strftime("%d %b %y")
            except ValueError:
                continue
    return ""


def _time(value: Any) -> str:
    """'2:30 pm', '14:30' or '1430 hrs' → '1430 Hr' (Part 1 para 42.2: four figures on the 24-hour
    clock and 'Hr' — never 'H' or 'hr'); '' if it can't be read."""
    s = _clean(value).lower()
    m = re.search(r"\b(\d{1,2})(?:[:.](\d{2}))?\s*([ap])\.?m\b", s)
    if m:
        h, mi = int(m.group(1)) % 12 + (12 if m.group(3) == "p" else 0), int(m.group(2) or 0)
    else:
        m = re.search(r"\b(\d{1,2})[:.]?(\d{2})\b", s)
        if not m:
            return ""
        h, mi = int(m.group(1)), int(m.group(2))
    return f"{h:02d}{mi:02d} Hr" if h < 24 and mi < 60 else ""


def _roman(n: int) -> str:
    out = ""
    for v, s in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"),
                 (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        while n >= v:
            out, n = out + s, n - v
    return out


# ── Word building blocks ─────────────────────────────────────────────────────────────────────────
def _fmt(p, align=JUSTIFY, left=None, tabs: Sequence = (), keep=False):
    """Paragraphs are justified (Part 1 para 49.2) unless a layout element says otherwise."""
    f = p.paragraph_format
    p.alignment = align
    if left is not None:
        f.left_indent = left
    for t in tabs:
        f.tab_stops.add_tab_stop(t)
    if keep:
        f.keep_with_next = True
    return p


def _run(p, text, bold=False, underline=False):
    r = p.add_run(text)
    r.bold = bold or None
    r.underline = underline or None
    return r


def _centre(p, text, bold=True, underline=False, keep=False):
    _fmt(p, align=CENTRE, keep=keep)
    _run(p, text, bold, underline)
    return p


def _blank(container, keep=False):
    """"Two-line spacing" is two line feeds — one blank line (Part 1 para 7)."""
    p = container.add_paragraph()
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def _para(p, number, text, heading="", keep=False):
    """Number at the margin, a 0.5 in tab, an optional bold heading with a bold full stop, then the
    text; later lines return to the margin under the number (Part 1 paras 33.3, 49.3-49.4)."""
    _fmt(p, tabs=(TAB,), keep=keep)
    p.add_run(f"{number}.\t")
    if heading:
        _run(p, heading.rstrip(".") + ".", bold=True)
        p.add_run(" ")
    p.add_run(_stop(text))
    return p


def _sub(p, number, text, indent=TAB, stop=True):
    """Sub-paragraph: the number sits under the first letter of the paragraph text, then a 0.5 in
    tab; later lines align under the number (Part 1 paras 33.4.4-33.4.5; App B note 30)."""
    _fmt(p, left=indent, tabs=(Emu(indent + TAB),))
    p.add_run(f"{number}\t")
    p.add_run(_stop(text) if stop else text)
    return p


def _keep(row):
    for c in row.cells:
        for p in c.paragraphs:
            p.paragraph_format.keep_with_next = True
    return row


def _table(container, widths, indent=0):
    """Borderless, fixed-width table whose text starts exactly at the margin, as in AD's layout."""
    t = container.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for col, w in zip(t._tbl.tblGrid.gridCol_lst, widths):
        col.w = Emu(w)
    pr = t._tbl.tblPr
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), str(round(indent / 635)))         # EMU → twips
    ind.set(qn("w:type"), "dxa")
    pr.insert_element_before(ind, "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                             "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), "0")
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    pr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    return t


def _row(t, widths, header=False):
    row = t.add_row()
    for c, w in zip(row.cells, widths):
        c.width = Emu(w)
    if header:                                           # repeats on every page the table reaches
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return row


def _fld_char(p, kind):
    r = OxmlElement("w:r")
    c = OxmlElement("w:fldChar")
    c.set(qn("w:fldCharType"), kind)
    r.append(c)
    p._p.append(r)


def _field(p, parts, shown):
    """A Word field: `parts` is its instruction as text and nested (parts, shown) fields; `shown` is
    displayed until Word recalculates it, which it does for header fields whenever it lays out pages."""
    _fld_char(p, "begin")
    for part in parts:
        if isinstance(part, str):
            r = OxmlElement("w:r")
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = part
            r.append(t)
            p._p.append(r)
        else:
            _field(p, *part)
    _fld_char(p, "separate")
    p.add_run(shown)
    _fld_char(p, "end")


def _pages_text(n: int) -> str:
    return "(Only page)" if n == 1 else f"({_WORDS[n].capitalize()} pages)" if n < 10 else f"({n} pages)"


def _page_count(p, estimate: int):
    """'(Only page)', '(Six pages)', '(27 pages)': words below ten, numerals from ten (Part 1 para
    18; App B note 2). Only Word knows the final page count, so Word computes it."""
    shown = _pages_text(estimate)
    n = ([" NUMPAGES "], str(estimate))
    words = ([" NUMPAGES \\* CardText \\* FirstCap "], _WORDS[min(estimate, 9)].capitalize())
    inner = ([" IF ", n, ' < 10 "(', words, ' pages)" "(', n, ' pages)" '], shown)
    _field(p, [" IF ", n, ' = 1 "(Only page)" "', inner, '" '], shown)


_SHAPETYPE = (
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e"><v:formulas>'
    + "".join(f'<v:f eqn="{e}"/>' for e in (
        "sum #0 0 10800", "prod #0 2 1", "sum 21600 0 @1", "sum 0 0 @2", "sum 21600 0 @3", "if @0 @3 0",
        "if @0 21600 @1", "if @0 0 @2", "if @0 @4 21600", "mid @5 @6", "mid @8 @5", "mid @7 @8",
        "mid @6 @7", "sum @6 0 @5"))
    + '</v:formulas><v:path textpathok="t" o:connecttype="custom" '
    'o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/>'
    '</v:handles><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>')


def _watermark(header, text, n):
    """Classified documents carry the classification diagonally across every page, beneath the text,
    with the copy number for SECRET and above (Part 1 paras 12.1, 76)."""
    w = 460
    h = min(130, round(w * 3 / max(len(text), 1)))
    header.paragraphs[0]._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
        f'<w:pict>{_SHAPETYPE}<v:shape id="JSSDWatermark{n}" o:spid="_x0000_s{2049 + n}" '
        f'type="#_x0000_t136" style="position:absolute;margin-left:0;margin-top:0;width:{w}pt;'
        f'height:{h}pt;rotation:315;z-index:-251654144;mso-position-horizontal:center;'
        'mso-position-horizontal-relative:margin;mso-position-vertical:center;'
        'mso-position-vertical-relative:margin" o:allowincell="f" fillcolor="silver" stroked="f">'
        f'<v:fill opacity=".5"/><v:textpath style="font-family:&quot;{FONT}&quot;;font-size:1pt" '
        f'string={quoteattr(text)}/></v:shape></w:pict></w:r>'))


def _arial(rpr):
    """Arial for every script, and no theme font that would override it."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in list(fonts.attrib):
        if a.endswith("Theme") or a.endswith("theme"):
            del fonts.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), FONT)


def _page_setup(doc, draft: bool):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = EDGE
    sec.gutter = GUTTER
    sec.header_distance = sec.footer_distance = EDGE      # Part 1 App E notes 1 and 14
    # Printed on both sides (Part 1 para 6), so the gutter swaps sides on the reverse page (App E note 15).
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))
    mirror = OxmlElement("w:mirrorMargins")
    if zoom is not None:
        zoom.addnext(mirror)
    else:
        settings.insert(0, mirror)

    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")) if defaults is not None else None
    if rpr is not None:
        _arial(rpr)
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), "en-GB")                    # Part 1 para 50: Oxford spellings

    normal = doc.styles["Normal"]
    _arial(normal.element.get_or_add_rPr())
    normal.font.size = SIZE
    pf = normal.paragraph_format
    pf.line_spacing = DRAFT_LINE_SPACING if draft else LINE_SPACING
    pf.space_before = pf.space_after = Pt(0)
    pf.alignment = JUSTIFY
    pf.widow_control = True


# ── the document, top to bottom ──────────────────────────────────────────────────────────────────
class _Minutes:
    """Numbers paragraphs 1, 2, 3 … straight through the minutes (AD)."""

    def __init__(self, doc):
        self.doc, self.n = doc, 0

    def next(self) -> int:
        self.n += 1
        return self.n


def _estimate_pages(doc) -> int:
    """Rough page count of the finished body; the fallback when LibreOffice can't lay it out."""
    per_inch = 12.5                                      # justified Arial 12 pt, as measured
    lines = 0

    def height(p, inches):
        chars = sum(len(t.text or "") for t in p.iter(qn("w:t")))
        return max(1, math.ceil(chars / max(inches * per_inch, 1)))

    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            lines += height(el, TEXT_W / 914400)
        elif el.tag == qn("w:tbl"):
            widths = [int(c.get(qn("w:w"))) / 1440 for c in el.iter(qn("w:gridCol"))]
            for tr in el.iter(qn("w:tr")):
                lines += max([1] + [sum(height(p, w) for p in tc.iter(qn("w:p")))
                                    for tc, w in zip(tr.iter(qn("w:tc")), widths)])
    return max(1, math.ceil(lines / LINES_PER_PAGE))


def _laid_out_pages(data: bytes) -> int:
    """Pages as LibreOffice lays the minutes out, or 0 if it can't. The image carries LibreOffice for
    .doc translation; the flags are the ones utils/documents.py explains for running it in a container."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return 0
    with tempfile.TemporaryDirectory() as tmpdir:
        src = os.path.join(tmpdir, "minutes.docx")
        with open(src, "wb") as fh:
            fh.write(data)
        cmd = [soffice, f"-env:UserInstallation=file://{os.path.join(tmpdir, 'loprofile')}",
               "--headless", "--norestore", "--nolockcheck", "--nodefault",
               "--convert-to", "pdf", "--outdir", tmpdir, src]
        try:
            subprocess.run(cmd, capture_output=True, timeout=SOFFICE_TIMEOUT,
                           env={**os.environ, "HOME": tmpdir}, check=False)
            pdf = pypdfium2.PdfDocument(os.path.join(tmpdir, "minutes.pdf"))
            pages = len(pdf)
            pdf.close()
            return pages
        except Exception as e:                           # the estimate stands; never fail the minutes
            logger.warning(f"[DOCX] could not lay out the minutes to count pages: {e}")
            return 0


def _superscription(doc, meta) -> bool:
    """Telephone and precedence on one line, the copy number under the precedence, the originator's
    address, then the file reference and date — each block two-line spaced (AD; Part 1 paras 19-20).
    Only what `meta` supplies is written."""
    tele, prec = _clean(meta.get("telephone")), _clean(meta.get("precedence")).upper()
    copy_no, address, file_ref = _clean(meta.get("copy_no")), _lines(meta.get("address")), _clean(meta.get("file_ref"))
    blocks = []
    if tele or prec:
        blocks.append([(tele if not tele or tele.lower().startswith("tele") else f"Tele: {tele}", prec)])
    if copy_no:
        blocks.append([("", copy_no if copy_no.lower().startswith("copy") else f"Copy No {copy_no}")])
    if address:
        blocks.append([(a, "") for a in address])
    if file_ref:
        # The date of issue is left blank for the signatory to write in ink (Part 1 para 41).
        blocks.append([(f"{file_ref}\tdt {_date(meta.get('issue_date'))}".rstrip(), "")])
    for i, block in enumerate(blocks):
        if i:
            _blank(doc)
        for left, right in block:
            p = _fmt(doc.add_paragraph(), align=LEFT)
            if left:
                p.add_run(left)
            if right:
                p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
                p.add_run("\t")
                _run(p, right, bold=right == prec)        # precedence: capitals and bold (Part 1 para 20.2.1)
    return bool(blocks)


def _title(doc, mom, meta):
    """A centre heading in block capitals, bold, not underlined, giving the place, time, date and
    purpose of the meeting (AD note 2; Part 1 paras 22 and 51). A long title breaks between phrases
    so that each line reads on its own (Part 1 App E note 11)."""
    venue = _clean(meta.get("venue")) or _clean(mom.get("venue"))
    when = _time(meta.get("meeting_time") or mom.get("meeting_time"))
    day = _date(meta.get("meeting_date") or mom.get("meeting_date"))
    head = "MINUTES OF THE MEETING" + (" HELD" if venue or when or day else "") + (f" AT {venue}" if venue else "")
    phrases = [head]
    at = " ".join(x for x in (f"AT {when}" if when else "", f"ON {day}" if day else "") if x)
    if at:
        phrases.append(at)
    topic = _clean(mom.get("title"))
    if topic:
        phrases.append(f"TO DISCUSS {topic}")
    lines: List[str] = []
    for phrase in phrases:
        if lines and len(lines[-1]) + 1 + len(phrase) <= 52:
            lines[-1] += " " + phrase
        else:
            lines.append(phrase)
    p = _fmt(doc.add_paragraph(), align=CENTRE, keep=True)
    for i, line in enumerate(lines):
        r = _run(p, line.upper().rstrip("."), bold=True)
        if i < len(lines) - 1:
            r.add_break()


def _present(w: _Minutes, attendees):
    """Paragraph 1: rank and name, appointment, and Chairman or Secretary, in columns; the chairman
    first, the others in the order given, the secretary last (AD note 3; Ch 6 para 16.4)."""
    rows = []
    for a in attendees:
        name, role = (_clean(a.get("name")), _clean(a.get("role"))) if isinstance(a, dict) else (_clean(a), "")
        if not name:
            continue
        parts = [x.strip() for x in role.split(",") if x.strip()]
        label = ("Chairman" if any(_CHAIR.fullmatch(x) for x in parts)
                 else "Secretary" if any(_SECRETARY.fullmatch(x) for x in parts) else "")
        # The label has its own column, so it is not repeated as the appointment.
        appointment = ", ".join(x for x in parts if not (_CHAIR.fullmatch(x) or _SECRETARY.fullmatch(x)))
        rows.append((name.strip("[]").replace("_", " "), appointment, label))
    rows.sort(key=lambda r: {"Chairman": 0, "Secretary": 2}.get(r[2], 1))   # stable: keeps given order
    if not rows:
        rows = [("Not identified from the recording", "", "")]

    n = w.next()
    _para(w.doc.add_paragraph(), n, "The following were present:-", keep=True)
    widths = (TEXT_W - TAB - Inches(2.9), Inches(1.9), COL)
    t = _table(w.doc, widths, indent=TAB)
    for i, (name, appointment, label) in enumerate(rows, 1):
        _keep(_row(t, widths)) if i == 1 else _row(t, widths)
        cells = _row(t, widths).cells
        _sub(cells[0].paragraphs[0], f"{n}.{i}.", name, indent=0, stop=False)
        _fmt(cells[1].paragraphs[0], align=LEFT).add_run(appointment)
        _fmt(cells[2].paragraphs[0], align=LEFT).add_run(label)


def _introduction(w: _Minutes, mom):
    """INTRODUCTION (Ch 6 para 16.12): why the meeting was held, the agenda taken up and the gist of
    the discussion. A centre heading is in capitals and bold, two-line spaced above and below, not
    underlined (Part 1 paras 8.1, 33.1 and 51)."""
    purpose = _clean(mom.get("purpose"))
    agenda = [a for a in (_clean(x) for x in mom.get("agenda") or []) if a]
    summary = [s for s in (_clean(x) for x in re.split(r"\n\s*\n", str(mom.get("summary") or ""))) if s]
    blocks = ([(purpose, [])] if purpose else []) \
        + ([("The following agenda was taken up:-", agenda)] if agenda else []) \
        + [(s, []) for s in summary]
    if not blocks:
        return
    _blank(w.doc, keep=True)
    _centre(w.doc.add_paragraph(), "INTRODUCTION", keep=True)
    for text, subs in blocks:
        _blank(w.doc, keep=True)
        n = w.next()
        _para(w.doc.add_paragraph(), n, text, keep=bool(subs))
        for j, s in enumerate(subs, 1):
            _blank(w.doc, keep=j == 1)        # a page never ends on a lead-in line (Part 1 para 36)
            _sub(w.doc.add_paragraph(), f"{n}.{j}.", s)


def _items(mom, grade: str) -> List[Dict[str, Any]]:
    """The discussion as AD items: the points discussed, then each decision, then each task as a
    decision with its owner in the Action column (Ch 6 paras 16.8, 16.14, 16.16)."""
    points = [p for p in (_clean(x) for x in mom.get("key_points") or []) if p]
    figures = [f for f in (_clean(x) for x in mom.get("key_figures") or []) if f]
    decisions = [(d, "") for d in (_clean(x) for x in mom.get("decisions") or []) if d]
    for ai in mom.get("action_items") or []:
        if not isinstance(ai, dict) or not _clean(ai.get("task")):
            continue
        task, due = _clean(ai.get("task")).rstrip("."), _clean(ai.get("due"))
        due = _date(due) or due
        if due and due.lower() not in task.lower():
            task += f" {due}" if re.match(r"(by|before|within|on|in|till|until|end of)\b", due, re.I) else f" by {due}"
        decisions.append((task, _clean(ai.get("assigned_to"))))
    if not (points or figures or decisions):
        return []
    return [{"title": _clean(mom.get("title")) or "Discussion", "grade": grade or "UNCLASSIFIED",
             "points": points, "figures": figures, "decisions": decisions}]


def _items_table(w: _Minutes, items):
    """Items in three borderless columns — text, Action and Info (AD). Each item is headed
    'ITEM I – TITLE', centred, bold and underlined, with its own classification in brackets below,
    even when unclassified (AD notes 4-5; Ch 6 paras 16.7 and 16.13). Action and Info are endorsed
    against each decision (AD note 6)."""
    widths = (TEXT_W - 2 * COL, COL, COL)
    _blank(w.doc)
    t = _table(w.doc, widths)
    head = _keep(_row(t, widths, header=True)).cells
    for c, label in zip(head[1:], ("Action", "Info")):
        _centre(c.paragraphs[0], label)                    # column headings bold, not capitals (AD)
    for k, item in enumerate(items, 1):
        _keep(_row(t, widths))
        _centre(_keep(_row(t, widths)).cells[0].paragraphs[0], f"ITEM {_roman(k)} – {item['title'].upper()}",
                underline=True, keep=True)
        _centre(_keep(_row(t, widths)).cells[0].paragraphs[0], f"({item['grade']})", bold=False, keep=True)
        entries = [(None, p, "") for p in item["points"]]
        if item["figures"]:
            entries.append(("figures", "The following figures were quoted:-", ""))
        entries += [("Decision", d, owner) for d, owner in item["decisions"]]
        for i, (kind, text, owner) in enumerate(entries):
            if i:
                _row(t, widths)
            else:
                _keep(_row(t, widths))
            cells = _row(t, widths).cells
            n = w.next()
            p = _para(cells[0].paragraphs[0], n, text, heading="Decision" if kind == "Decision" else "",
                      keep=kind == "figures")
            p.paragraph_format.right_indent = Inches(0.1)   # clear of the Action column
            _centre(cells[1].paragraphs[0], owner, bold=False)
            if kind == "figures":
                for j, f in enumerate(item["figures"], 1):
                    _keep(_row(t, widths)) if j == 1 else _row(t, widths)
                    sub = _sub(_row(t, widths).cells[0].paragraphs[0], f"{n}.{j}.", f)
                    sub.paragraph_format.right_indent = Inches(0.1)


def _closing(w: _Minutes, meta):
    """The standard closing paragraph, then the secretary's signature block nine line feeds below it
    (AD; Ch 6 paras 15 and 16.17: the secretary signs once the chairman has approved the draft)."""
    _blank(w.doc)
    by = _date(meta.get("amendments_by")) or DOTS
    text = f"Agreement with the minutes will be assumed unless amendments are received by {by}"
    _para(w.doc.add_paragraph(), w.next(), text, keep=True)
    for _ in range(SIGNATURE_FEEDS - 1):
        _blank(w.doc, keep=True)
    sec = meta.get("secretary") if isinstance(meta.get("secretary"), dict) else {}
    name = _clean(sec.get("name") or meta.get("secretary_name")) or "Initials and Name"
    rank = _clean(sec.get("rank") or meta.get("secretary_rank")) or "Rank"
    for i, line in enumerate((f"({name})", rank, "Secretary")):   # Part 1 para 27.1
        _fmt(w.doc.add_paragraph(), align=LEFT, keep=i < 2).add_run(line)


def _distribution(doc, meta):
    """Distribution after the signature block: Distribution, No of Copies, Copy No and Remarks, the
    headings bold and not underlined (AD; Part 1 paras 51 and 67-69). The office file copy is always
    listed, as in every JSSD specimen."""
    rows = []
    for d in meta.get("distribution") or []:
        if isinstance(d, dict) and _clean(d.get("addressee")):
            rows.append((_clean(d.get("addressee")), _clean(d.get("copies")) or "One",
                         _clean(d.get("copy_no")), _clean(d.get("remarks"))))
        elif isinstance(d, str) and _clean(d):
            rows.append((_clean(d), "One", "", ""))
    if not any(r[0].lower() == "file" for r in rows):
        rows.append(("File", "One", "", ""))
    _blank(doc, keep=True)
    widths = (Inches(2.4), Inches(1.4), Inches(1.1), TEXT_W - Inches(4.9))
    t = _table(doc, widths)
    head = _keep(_row(t, widths)).cells
    for c, label in zip(head, ("Distribution", "No of Copies", "Copy No", "Remarks")):
        _run(_fmt(c.paragraphs[0], align=LEFT), label, bold=True)
    for r in rows:
        _row(t, widths)
        for c, v in zip(_row(t, widths).cells, r):
            _fmt(c.paragraphs[0], align=LEFT).add_run(v)


def _headers(doc, grade: str, copy_no: str, pages: int):
    """Classification centred in bold capitals at the head and foot of every page, 0.5 in from the
    edge, two-line spaced from the text (Part 1 paras 12.1, App B notes 27-28). Page 1 adds the number
    of pages one line below it and no page number; later pages carry the page number two lines below
    it (Part 1 paras 16 and 18). Unclassified documents carry no marking at all (Part 1 para 11.5)."""
    normal = doc.styles["Normal"]
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    first, rest = sec.first_page_header, sec.header

    def para(part, i):
        p = part.paragraphs[0] if i == 0 else part.add_paragraph()
        p.style = normal
        return p

    def page_number(p):
        _fmt(p, align=CENTRE)
        _field(p, [" PAGE "], "2")

    if not grade:
        p = para(first, 0)                                 # nothing above the first line of text
        p.paragraph_format.line_spacing = Pt(1)
        page_number(para(rest, 0))
        para(rest, 1)
        return

    _centre(para(first, 0), grade)
    if grade in _PAGES_SHOWN:
        p = _fmt(para(first, 1), align=CENTRE)
        _page_count(p, pages)
    para(first, 2)
    _centre(para(rest, 0), grade)
    para(rest, 1)
    page_number(para(rest, 2))
    para(rest, 3)
    for footer in (sec.first_page_footer, sec.footer):
        para(footer, 0)
        _centre(para(footer, 1), grade)
    mark = grade
    if copy_no and grade in _COPY_IN_WATERMARK:
        mark += "  " + (copy_no if copy_no.lower().startswith("copy") else f"Copy No {copy_no}")
    _watermark(first, mark, 1)
    _watermark(rest, mark, 2)


def build_mom_docx(mom: Dict[str, Any], meta: Dict[str, Any] = None) -> bytes:
    """MoM response object, plus the job's optional `mom_meta`, → a JSSD-format .docx, ready to upload.

    Page 1 of CONFIDENTIAL and higher minutes states the number of pages. That is a Word field, which
    Word recalculates on opening, but other viewers show the number stored with it — so the stored
    number is the count LibreOffice lays out, falling back to an estimate."""
    meta = meta if isinstance(meta, dict) else {}
    data, shown = _build(mom, meta)
    if _grade(meta.get("classification")) in _PAGES_SHOWN:
        laid_out = _laid_out_pages(data)
        if laid_out and laid_out != shown:
            data, _ = _build(mom, meta, laid_out)
    return data


def _build(mom: Dict[str, Any], meta: Dict[str, Any], pages: int = 0) -> Tuple[bytes, int]:
    grade, draft = _grade(meta.get("classification")), bool(meta.get("draft"))
    doc = docx.Document()
    _page_setup(doc, draft)
    w = _Minutes(doc)

    if draft:   # Part 1 para 73: 'DRAFT', centred, two lines below the classification
        _centre(doc.add_paragraph(), "DRAFT")
        _blank(doc)
    if _superscription(doc, meta):
        _blank(doc)
    _title(doc, mom, meta)
    _blank(doc, keep=True)
    _present(w, mom.get("attendees") or [])
    _introduction(w, mom)
    items = _items(mom, grade)
    if items:
        _items_table(w, items)
    _closing(w, meta)
    _distribution(doc, meta)
    pages = pages or _estimate_pages(doc)
    _headers(doc, grade, _clean(meta.get("copy_no")), pages)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), pages


# ── translation output ────────────────────────────────────────────────────────────────────────────
# Not the JSSD layout: that governs minutes of a meeting, and a translated recording is not minutes.
# A plain document instead, in the same Arial and spacing so the two outputs look like one product:
# the translation first, because that is what was asked for, then the original transcript under it
# so a reader can check any line against what was actually said.

def build_translation_docx(result: Dict[str, Any], file_name: str = "") -> bytes:
    """/translate-media response → .docx bytes: title, languages, translation, original transcript."""
    doc = docx.Document()
    _page_setup(doc, draft=False)

    source = result.get("source_lang") or ""
    target = result.get("target_lang") or ""
    title = f"TRANSLATION OF {file_name.upper()}" if file_name else "TRANSLATION"
    _centre(doc.add_paragraph(), title, keep=True)
    _blank(doc)
    meta = f"{source} to {target}"
    if result.get("duration_s"):
        minutes, seconds = divmod(int(result["duration_s"]), 60)
        meta += f", {minutes} min {seconds:02d} s of audio"
    _centre(doc.add_paragraph(), meta, bold=False)

    for heading, text in ((f"TRANSLATION ({target.upper()})", result.get("translated_text", "")),
                          (f"ORIGINAL TRANSCRIPT ({source.upper()})", result.get("original_text", ""))):
        _blank(doc, keep=True)
        _centre(doc.add_paragraph(), heading, keep=True)
        for block in [b.strip() for b in (text or "").split("\n\n") if b.strip()]:
            _blank(doc)
            _fmt(doc.add_paragraph()).add_run(block)

    notes = result.get("notes") or []
    if notes:
        _blank(doc, keep=True)
        _centre(doc.add_paragraph(), "NOTES", keep=True)
        for note in notes:
            _blank(doc)
            _fmt(doc.add_paragraph()).add_run(str(note))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
